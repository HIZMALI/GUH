"""Read-only source extraction; never changes supplied PDFs/workbook."""
from pathlib import Path
import hashlib
import json
from datetime import date, datetime, time
from pypdf import PdfReader
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'data' / 'source'
OUT.mkdir(parents=True, exist_ok=True)
manifest = []
for path in sorted(ROOT.glob('*.pdf')):
    reader = PdfReader(path)
    pages = [page.extract_text(extraction_mode='layout') or '' for page in reader.pages]
    (OUT / (path.stem + '.txt')).write_text('\n\n'.join(f'--- PAGE {i+1} ---\n{p}' for i,p in enumerate(pages)), encoding='utf-8')
    manifest.append({'file':path.name,'sha256':hashlib.sha256(path.read_bytes()).hexdigest(),'pages':len(pages),'characters':[len(p) for p in pages]})
for path in ROOT.glob('*.xlsx'):
    book = load_workbook(path, read_only=True, data_only=True)
    sheets = {ws.title: [[v.isoformat() if isinstance(v,(date,datetime,time)) else v for v in row] for row in ws.iter_rows(values_only=True)] for ws in book}
    (OUT / 'workbook.json').write_text(json.dumps(sheets,ensure_ascii=False,indent=2),encoding='utf-8')
    manifest.append({'file':path.name,'sha256':hashlib.sha256(path.read_bytes()).hexdigest(),'sheets':{k:{'rows':len(v),'columns':max(map(len,v),default=0)} for k,v in sheets.items()}})
    book.close()
(OUT / 'manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(manifest,ensure_ascii=False,indent=2))
