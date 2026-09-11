"""Normalize the organizer workbook without inventing missing sensor channels."""
from pathlib import Path
from datetime import datetime, time
import hashlib
import json
import math
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]


def parse_workbook(path: Path) -> dict:
    values = load_workbook(path, read_only=True, data_only=True)
    formulas = load_workbook(path, read_only=True, data_only=False)
    rows = list(values['Akım Sensörü'].iter_rows(values_only=True))
    samples = []
    previous = None
    day = 0
    for index, row in enumerate(rows, 1):
        stamp, secondary, unit, ratio, primary, primary_unit = row
        if not isinstance(secondary, (float,int)) or not isinstance(primary, (float,int)):
            continue
        if unit != 'mA' or primary_unit != 'A':
            raise ValueError(f'Unexpected unit at row {index}')
        if not isinstance(stamp, (datetime,time)):
            raise ValueError(f'Unexpected time at row {index}')
        seconds = stamp.hour * 3600 + stamp.minute * 60 + stamp.second
        if previous is not None and seconds < previous:
            day += 1
        offset = day * 86400 + seconds
        computed = secondary * ratio / 1000
        if not math.isclose(computed, primary, abs_tol=1e-8):
            raise ValueError(f'Cached formula disagrees with B*D/1000 at row {index}')
        formula = formulas['Akım Sensörü'].cell(index,5).value
        if formula != f'=B{index}*D{index}/1000':
            raise ValueError(f'Unexpected source formula at E{index}')
        if samples and offset - samples[-1]['offset_seconds'] != 900:
            raise ValueError('Nonuniform sample interval requires explicit handling')
        samples.append({'source_row':index,'source_cell':f'E{index}','offset_seconds':offset,'original_time':stamp.isoformat(),'secondary_ma':secondary,'ratio':ratio,'current_l1':primary,'source':'organizer_synthetic_replay'})
        previous = seconds
    values.close()
    formulas.close()
    return {'source_file':path.name,'source_sha256':hashlib.sha256(path.read_bytes()).hexdigest(),'sheet':'Akım Sensörü','sample_interval_seconds':900,'original_timestamps_are_elapsed':True,'measurements_available':['current_l1'],'missing_channels':['current_l2','current_l3','voltage','temperature','humidity','pd','arc'],'samples':samples}


if __name__ == '__main__':
    output = ROOT / 'data/synthetic/current_replay.json'
    output.parent.mkdir(parents=True,exist_ok=True)
    dataset = parse_workbook(ROOT / 'İstenen Veriler.xlsx')
    output.write_text(json.dumps(dataset,ensure_ascii=False,indent=2),encoding='utf-8')
    print(f'Validated {len(dataset["samples"])} original L1 rows -> {output.relative_to(ROOT)}')
